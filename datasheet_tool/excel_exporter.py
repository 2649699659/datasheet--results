"""
excel_exporter.py — Step 4: Export to Excel

Simple 3-sheet Excel output:
- All Parameters: All extracted parameters
- Needs Review: Only low-confidence/unresolved parameters
- Document Info: Metadata about the document

Input:  ExtractionResult (parameters + metadata)
Output: Excel file
"""

import logging
from typing import List

from .models import Parameter, DocumentInfo, ExtractionResult

logger = logging.getLogger(__name__)


def _auto_adjust_column_width(ws, min_width: int = 8, max_width: int = 50):
    """Auto-adjust column widths based on content."""
    try:
        from openpyxl.utils import get_column_letter
    except ImportError:
        return
    
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        
        adjusted_width = min(max(max_length + 2, min_width), max_width)
        ws.column_dimensions[column_letter].width = adjusted_width


def _write_metadata_sheet(wb, doc_info: DocumentInfo):
    """Write the Document Info sheet."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        logger.error("openpyxl not available")
        return
    
    ws = wb.active
    ws.title = "Document Info"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Headers
    headers = ["Field", "Value"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    
    # Data
    fields = [
        ("Manufacturer", doc_info.manufacturer),
        ("Part Number", doc_info.part_number),
        ("Module Type", doc_info.module_type),
        ("Description", doc_info.description),
        ("File Name", doc_info.file_name),
        ("Processing Time (s)", f"{doc_info.processing_time_seconds:.1f}"),
        ("Total Parameters", doc_info.total_parameters),
        ("Needs Review", doc_info.needs_review_count),
    ]
    
    for row_idx, (field, value) in enumerate(fields, start=2):
        ws.cell(row=row_idx, column=1, value=field)
        ws.cell(row=row_idx, column=2, value=value or "")
    
    _auto_adjust_column_width(ws)


def _write_all_parameters_sheet(wb, parameters: List[Parameter]):
    """Write the All Parameters sheet."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        logger.error("openpyxl not available")
        return
    
    ws = wb.create_sheet("All Parameters")
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Headers
    headers = [
        "Page", "Table", "Row", "Symbol", "Name",
        "Value", "Min", "Typ", "Max", "Unit",
        "Condition", "Confidence", "Needs Review", "Source Text"
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    
    # Data
    for row_idx, param in enumerate(parameters, start=2):
        ws.cell(row=row_idx, column=1, value=param.page)
        ws.cell(row=row_idx, column=2, value=param.table_index)
        ws.cell(row=row_idx, column=3, value=param.row_index)
        ws.cell(row=row_idx, column=4, value=param.symbol or "")
        ws.cell(row=row_idx, column=5, value=param.name or "")
        
        # Value fields
        ws.cell(row=row_idx, column=6, value=str(param.value) if param.value else "")
        ws.cell(row=row_idx, column=7, value=str(param.min) if param.min else "")
        ws.cell(row=row_idx, column=8, value=str(param.typ) if param.typ else "")
        ws.cell(row=row_idx, column=9, value=str(param.max) if param.max else "")
        
        ws.cell(row=row_idx, column=10, value=param.unit or "")
        ws.cell(row=row_idx, column=11, value=param.condition or "")
        ws.cell(row=row_idx, column=12, value=round(param.confidence, 2) if param.confidence else 1.0)
        ws.cell(row=row_idx, column=13, value="Yes" if param.needs_review else "No")
        ws.cell(row=row_idx, column=14, value=param.source_text or "")
    
    _auto_adjust_column_width(ws)


def _write_needs_review_sheet(wb, parameters: List[Parameter]):
    """Write the Needs Review sheet (only low-confidence/unresolved parameters)."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        logger.error("openpyxl not available")
        return
    
    ws = wb.create_sheet("Needs Review")
    
    # Filter to only needs_review
    review_params = [p for p in parameters if p.should_review()]
    
    if not review_params:
        ws.cell(row=1, column=1, value="No parameters need review")
        return
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="FFC000")  # Orange
    center_align = Alignment(horizontal="center", vertical="center")
    
    # Headers
    headers = [
        "Page", "Table", "Row", "Symbol", "Name",
        "Value", "Min", "Typ", "Max", "Unit",
        "Condition", "Confidence", "Issue", "Source Text"
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
    
    # Data
    for row_idx, param in enumerate(review_params, start=2):
        ws.cell(row=row_idx, column=1, value=param.page)
        ws.cell(row=row_idx, column=2, value=param.table_index)
        ws.cell(row=row_idx, column=3, value=param.row_index)
        ws.cell(row=row_idx, column=4, value=param.symbol or "")
        ws.cell(row=row_idx, column=5, value=param.name or "")
        
        ws.cell(row=row_idx, column=6, value=str(param.value) if param.value else "")
        ws.cell(row=row_idx, column=7, value=str(param.min) if param.min else "")
        ws.cell(row=row_idx, column=8, value=str(param.typ) if param.typ else "")
        ws.cell(row=row_idx, column=9, value=str(param.max) if param.max else "")
        
        ws.cell(row=row_idx, column=10, value=param.unit or "")
        ws.cell(row=row_idx, column=11, value=param.condition or "")
        ws.cell(row=row_idx, column=12, value=round(param.confidence, 2) if param.confidence else 1.0)
        
        # Issue description
        issues = []
        if param.confidence < 0.8:
            issues.append(f"Low confidence ({param.confidence:.2f})")
        if not param.unit and (param.value or param.min or param.typ or param.max):
            issues.append("Missing unit")
        if param.unit and not (param.value or param.min or param.typ or param.max):
            issues.append("Missing value")
        if not (param.value or param.min or param.typ or param.max):
            issues.append("No value")
        ws.cell(row=row_idx, column=13, value="; ".join(issues))
        
        ws.cell(row=row_idx, column=14, value=param.source_text or "")
    
    _auto_adjust_column_width(ws)


def export(result: ExtractionResult, output_path: str):
    """
    Export extraction result to Excel file.

    Args:
        result: ExtractionResult with parameters and metadata
        output_path: Path to output Excel file
    """
    logger.info(f"Writing Excel to: {output_path}")
    
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl not installed. Run: pip install openpyxl")
        raise
    
    wb = openpyxl.Workbook()
    
    # Write sheets
    _write_metadata_sheet(wb, result.document_info)
    _write_all_parameters_sheet(wb, result.parameters)
    _write_needs_review_sheet(wb, result.parameters)
    
    # Save
    wb.save(output_path)
    logger.info(f"Excel saved: {output_path}")


def export_from_dict(result_dict: dict, output_path: str):
    """
    Export from a dict (for compatibility with simple_extractor output).
    
    Args:
        result_dict: Dict with "metadata" and "parameters" keys
        output_path: Path to output Excel file
    """
    # Convert dict to ExtractionResult
    doc_info = DocumentInfo(
        manufacturer=result_dict.get("metadata", {}).get("manufacturer"),
        part_number=result_dict.get("metadata", {}).get("part_number"),
        module_type=result_dict.get("metadata", {}).get("module_type"),
        description=result_dict.get("metadata", {}).get("description"),
        file_name=result_dict.get("metadata", {}).get("file_name", ""),
        pdf_path=result_dict.get("metadata", {}).get("pdf_path", ""),
        processing_time_seconds=result_dict.get("metadata", {}).get("processing_time_seconds", 0),
        total_parameters=len(result_dict.get("parameters", [])),
        needs_review_count=sum(
            1 for p in result_dict.get("parameters", [])
            if p.get("needs_review") or _check_needs_review(p)
        ),
    )
    
    parameters = []
    for p_dict in result_dict.get("parameters", []):
        param = Parameter(
            symbol=p_dict.get("symbol"),
            name=p_dict.get("name"),
            value=p_dict.get("value"),
            min=p_dict.get("min"),
            typ=p_dict.get("typ"),
            max=p_dict.get("max"),
            unit=p_dict.get("unit"),
            condition=p_dict.get("condition"),
            page=p_dict.get("page", 0),
            table_index=p_dict.get("table_index", 0),
            row_index=p_dict.get("row_index", 0),
            source_text=p_dict.get("source_text"),
            confidence=p_dict.get("confidence", 1.0),
            needs_review=p_dict.get("needs_review", False),
        )
        parameters.append(param)
    
    result = ExtractionResult(document_info=doc_info, parameters=parameters)
    export(result, output_path)


def _check_needs_review(p: dict) -> bool:
    """Check if a parameter dict needs review."""
    confidence = p.get("confidence", 1.0)
    if confidence < 0.8:
        return True
    has_value = p.get("value") or p.get("min") or p.get("typ") or p.get("max")
    has_unit = p.get("unit")
    if has_value and not has_unit:
        return True
    if has_unit and not has_value:
        return True
    if not has_value:
        return True
    return False


if __name__ == "__main__":
    import sys
    import json
    
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    
    if len(sys.argv) < 3:
        print("Usage: python3 excel_exporter.py <json_file> <output_xlsx>")
        sys.exit(1)
    
    json_path = sys.argv[1]
    output_path = sys.argv[2]
    
    with open(json_path, "r", encoding="utf-8") as f:
        result_dict = json.load(f)
    
    export_from_dict(result_dict, output_path)
    print(f"Excel exported: {output_path}")
